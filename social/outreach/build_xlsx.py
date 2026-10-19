"""Build x_accounts_vN.xlsx + .csv from a TSV.
Usage: python3 build_xlsx.py input.tsv x_accounts_v2
TSV columns (with header row): Категория, Когорта, Аккаунт, Имя, Подписчики, Тип, Описание, Почему полезен, Тегать/Писать, Ссылка
(For v1 the Когорта / Тегать/Писать columns may be empty.)"""
import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
H = ["Категория","Когорта","Аккаунт","Имя","Подписчики","Тип","Описание","Почему полезен","Тегать/Писать","Ссылка"]
src, out = sys.argv[1], sys.argv[2]
rows = list(csv.DictReader(open(src, encoding="utf-8"), delimiter="\t"))
for r in rows:
    r["Подписчики"] = int(str(r.get("Подписчики","0")).replace(",","").replace(" ","") or 0)
    if not r.get("Ссылка"): r["Ссылка"] = "https://x.com/" + r["Аккаунт"].lstrip("@")
order = {"large":0,"medium":1,"small":2,"sub-1K":3}
rows.sort(key=lambda r: (r["Категория"], order.get(r.get("Когорта",""), 9), -r["Подписчики"]))
wb = Workbook(); wb.remove(wb.active)
def sheet(name, data):
    ws = wb.create_sheet(name[:31]); ws.append(H)
    for c in ws[1]: c.font = Font(bold=True)
    for r in data: ws.append([r.get(h,"") for h in H])
    for i,w in enumerate([18,10,22,24,12,12,50,50,14,32],1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
sheet("All", rows)
for cat in dict.fromkeys(r["Категория"] for r in rows):
    sheet(cat.replace("/","-"), [r for r in rows if r["Категория"] == cat])
wb.save(out + ".xlsx")
with open(out + ".csv", "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=H, extrasaction="ignore"); w.writeheader(); w.writerows(rows)
print(len(rows), "rows ->", out + ".xlsx/.csv")
